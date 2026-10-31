"""Word tracked-changes engine for NDA Review.

Implements NCP's locked redline workflow natively (Streamlit Cloud has no
LibreOffice):

  1. Baseline: the counterparty's file with any tracked changes they sent
     accepted, so NCP's markup is a single layer.
  2. NCP's position: revised text for the paragraphs NCP changes, plus any
     new clauses, keyed to baseline paragraph numbers.
  3. Redline by comparison: each changed paragraph is diffed word by word
     against the baseline and the differences are written as w:ins / w:del
     inside the counterparty's own runs, so their formatting, numbering,
     hyperlinks and layout survive. Paragraphs NCP does not change are left
     exactly as the counterparty sent them.
  4. Audit before delivery: reject-all must reproduce the baseline,
     accept-all must equal NCP's position, untouched paragraphs must be
     unchanged, and the revision markup must be well formed, authored
     "New Capital Partners", with IDs from 960.

The CLEAN deliverable is the redline with every change accepted.
"""

import copy
import difflib
import io
import re
from datetime import datetime, timezone

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = "{%s}" % W_NS
MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"

AUTHOR = "New Capital Partners"
ID_START = 960
# Shown to the reviewer in place of a line break inside a paragraph, so
# each paragraph stays on one line of the document view.
LINE_BREAK_MARK = "⏎"

_CONTAINERS = {W + t for t in ("hyperlink", "smartTag", "customXml", "sdt",
                               "sdtContent", "fldSimple", "dir", "bdo")}
_INS = {W + "ins", W + "moveTo"}
_DEL = {W + "del", W + "moveFrom"}
_MARKER_PARENTS = {W + "rPr", W + "trPr", W + "numPr", W + "tcPr"}
_PROPERTY_CHANGES = ("rPrChange", "pPrChange", "sectPrChange", "tblPrChange",
                     "tblGridChange", "trPrChange", "tcPrChange",
                     "tblPrExChange", "numberingChange")
_RANGE_MARKS = ("moveFromRangeStart", "moveFromRangeEnd", "moveToRangeStart",
                "moveToRangeEnd", "customXmlInsRangeStart",
                "customXmlInsRangeEnd", "customXmlDelRangeStart",
                "customXmlDelRangeEnd")
_STORY_PART = re.compile(r"^/word/(header|footer)\d*\.xml$|"
                         r"^/word/(footnotes|endnotes)\.xml$")
SIGNATURE_CUE = re.compile(
    r"^\s*(in witness whereof|dated\b|date:|by:|by\s*_|name:|title:|"
    r"signature|authorized signature|accepted and agreed|agreed and accepted|"
    r"acknowledged and agreed|agreed to and accepted|very truly yours|"
    r"sincerely|\[signature)", re.I)


# ── Package helpers ──────────────────────────────────────────────────

def _load(docx_bytes):
    import docx
    return docx.Document(io.BytesIO(docx_bytes))


def _save(doc):
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _el(tag):
    from docx.oxml import OxmlElement
    return OxmlElement("w:" + tag)


def _is_tag(el):
    return isinstance(el.tag, str)


def _unwrap(el):
    parent = el.getparent()
    idx = parent.index(el)
    for child in list(el):
        parent.insert(idx, child)
        idx += 1
    parent.remove(el)


def _in_textbox(p):
    a = p.getparent()
    while a is not None:
        if a.tag == W + "txbxContent" or a.tag == MC + "AlternateContent":
            return True
        if a.tag == W + "body":
            return False
        a = a.getparent()
    return False


def list_paragraphs(doc):
    """Body paragraphs in document order (table cells included, text-box
    content excluded)."""
    return [p for p in doc.element.body.iter(W + "p") if not _in_textbox(p)]


# ── Text views ───────────────────────────────────────────────────────

def _run_atoms(r):
    """[(child, visible_text)] for a run's content (rPr excluded)."""
    atoms = []
    for ch in r:
        if not _is_tag(ch) or ch.tag == W + "rPr":
            continue
        tag = ch.tag
        if tag in (W + "t", W + "delText"):
            atoms.append((ch, ch.text or ""))
        elif tag in (W + "tab", W + "ptab"):
            atoms.append((ch, "\t"))
        elif tag in (W + "br", W + "cr"):
            atoms.append((ch, "\n"))
        elif tag == W + "noBreakHyphen":
            atoms.append((ch, "‑"))
        else:
            atoms.append((ch, ""))
    return atoms


def _run_text(r):
    return "".join(t for _, t in _run_atoms(r))


def para_text(p, mode="accept"):
    """Visible text of one paragraph. mode: 'accept' (insertions shown,
    deletions hidden) or 'reject' (the reverse)."""
    parts = []

    def walk(el):
        for ch in el:
            if not _is_tag(ch):
                continue
            tag = ch.tag
            if tag == W + "r":
                parts.append(_run_text(ch))
            elif tag in _INS:
                if mode != "reject":
                    walk(ch)
            elif tag in _DEL:
                if mode == "reject":
                    walk(ch)
            elif tag in _CONTAINERS:
                walk(ch)
    walk(p)
    return "".join(parts)


def _mark(p):
    rpr = p.find(W + "pPr/" + W + "rPr")
    if rpr is None:
        return None
    if rpr.find(W + "ins") is not None or rpr.find(W + "moveTo") is not None:
        return "ins"
    if rpr.find(W + "del") is not None or rpr.find(W + "moveFrom") is not None:
        return "del"
    return None


def doc_view(paras, mode):
    """Paragraph texts as Word shows them after accept-all or reject-all,
    honoring paragraph-mark revisions (a deleted mark joins a paragraph to
    the next one)."""
    texts, carry = [], ""
    for p in paras:
        t = carry + para_text(p, mode)
        m = _mark(p)
        if (mode == "accept" and m == "del") or (mode == "reject" and m == "ins"):
            carry = t
            continue
        texts.append(t)
        carry = ""
    if carry:
        texts.append(carry)
    return texts


def count_revisions(docx_bytes):
    """(insertions, deletions) carried by the file's body."""
    doc = _load(docx_bytes)
    body = doc.element.body
    return (sum(1 for _ in body.iter(W + "ins", W + "moveTo")),
            sum(1 for _ in body.iter(W + "del", W + "moveFrom")))


# ── Accept all (baseline + clean deliverable) ───────────────────────

def _merge_into_next(p):
    nxt = p.getnext()
    while nxt is not None and not _is_tag(nxt):
        nxt = nxt.getnext()
    if nxt is None or nxt.tag != W + "p":
        return False
    content = [c for c in p if _is_tag(c) and c.tag != W + "pPr"]
    idx = 1 if (len(nxt) and nxt[0].tag == W + "pPr") else 0
    for c in content:
        nxt.insert(idx, c)
        idx += 1
    p.getparent().remove(p)
    return True


def _accept_tree(root):
    for tag in _RANGE_MARKS:
        for el in list(root.iter(W + tag)):
            el.getparent().remove(el)
    for tag in _PROPERTY_CHANGES:
        for el in list(root.iter(W + tag)):
            el.getparent().remove(el)
    for el in list(root.iter(W + "del", W + "moveFrom")):
        parent = el.getparent()
        if parent is not None and parent.tag not in _MARKER_PARENTS:
            parent.remove(el)
    for el in list(root.iter(W + "ins", W + "moveTo")):
        parent = el.getparent()
        if parent is not None and parent.tag not in _MARKER_PARENTS:
            _unwrap(el)
    for el in list(root.iter(W + "cellIns", W + "cellDel")):
        el.getparent().remove(el)
    # Revision markers. Paragraph marks are processed last-first so a run
    # of deleted marks collapses into the surviving paragraph.
    markers = [el for el in root.iter(W + "ins", W + "del", W + "moveFrom",
                                      W + "moveTo")
               if el.getparent() is not None
               and el.getparent().tag in _MARKER_PARENTS]
    for el in reversed(markers):
        holder = el.getparent()
        if holder is None:
            continue
        owner = holder.getparent()
        deleted = el.tag in (W + "del", W + "moveFrom")
        holder.remove(el)
        if not deleted or owner is None:
            continue
        if holder.tag == W + "trPr":
            row = owner
            if row.getparent() is not None:
                row.getparent().remove(row)
        elif holder.tag == W + "rPr" and owner.tag == W + "pPr":
            p = owner.getparent()
            if p is not None and p.tag == W + "p":
                _merge_into_next(p)


def accept_all(docx_bytes):
    """Accept every tracked change (body, headers, footers, notes)."""
    from lxml import etree
    doc = _load(docx_bytes)
    _accept_tree(doc.element.body)
    for part in doc.part.package.iter_parts():
        if not _STORY_PART.match(str(part.partname)):
            continue
        if hasattr(part, "_element"):
            _accept_tree(part._element)
        else:
            # python-docx loads footnotes/endnotes as opaque parts.
            root = etree.fromstring(part.blob)
            _accept_tree(root)
            part._blob = etree.tostring(root, xml_declaration=True,
                                        encoding="UTF-8", standalone=True)
    return _save(doc)


_REVISION_TAG = re.compile(rb"<w:(ins|del|moveFrom|moveTo)[\s/>]")


def story_revisions(docx_bytes):
    """Live revision marks in every story part (body, headers, footers,
    notes), read straight from the package."""
    import zipfile
    total = 0
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as z:
        for name in z.namelist():
            if re.match(r"^word/(document|header\d*|footer\d*|footnotes|"
                        r"endnotes)\.xml$", name):
                total += len(_REVISION_TAG.findall(z.read(name)))
    return total


# ── Numbering labels (so the reviewer can cite clause numbers) ───────

def _roman(n):
    vals = [(1000, "m"), (900, "cm"), (500, "d"), (400, "cd"), (100, "c"),
            (90, "xc"), (50, "l"), (40, "xl"), (10, "x"), (9, "ix"),
            (5, "v"), (4, "iv"), (1, "i")]
    out = ""
    for v, s in vals:
        while n >= v:
            out += s
            n -= v
    return out


def _fmt_num(n, fmt):
    if fmt in ("lowerLetter", "upperLetter"):
        s, k = "", n
        while k > 0:
            k, rem = divmod(k - 1, 26)
            s = chr(97 + rem) + s
        return s.upper() if fmt == "upperLetter" else s
    if fmt in ("lowerRoman", "upperRoman"):
        s = _roman(n)
        return s.upper() if fmt == "upperRoman" else s
    if fmt == "bullet":
        return "•"
    if fmt == "none":
        return ""
    return str(n)


def _val(el, path):
    if el is None:
        return None
    x = el.find(path)
    return None if x is None else x.get(W + "val")


def _lvl_def(lvl):
    return (_val(lvl, W + "numFmt") or "decimal",
            _val(lvl, W + "lvlText") or "",
            int(_val(lvl, W + "start") or "1"),
            lvl.find(W + "isLgl") is not None)


def numbering_info(doc, paras):
    """Best-effort rendering of Word auto-numbering: [(label, ilvl)] per
    paragraph, with ilvl None for paragraphs that are not numbered."""
    try:
        numbering = doc.part.numbering_part.element
    except Exception:
        numbering = None
    try:
        styles = {s.get(W + "styleId"): s
                  for s in doc.styles.element.iter(W + "style")}
    except Exception:
        styles = {}
    if numbering is None:
        return [("", None)] * len(paras)
    abstract, links, defines, nums = {}, {}, {}, {}
    for a in numbering.iter(W + "abstractNum"):
        aid = a.get(W + "abstractNumId")
        abstract[aid] = {int(l.get(W + "ilvl", "0")): _lvl_def(l)
                         for l in a.iter(W + "lvl")}
        link = _val(a, W + "numStyleLink")
        if link:
            links[aid] = link
        defined = _val(a, W + "styleLink")
        if defined:
            defines[defined] = aid
    for n in numbering.iter(W + "num"):
        nid = n.get(W + "numId")
        restarts, lvl_over = {}, {}
        for ov in n.iter(W + "lvlOverride"):
            k = int(ov.get(W + "ilvl", "0"))
            so = _val(ov, W + "startOverride")
            if so is not None:
                restarts[k] = int(so)
            lv = ov.find(W + "lvl")
            if lv is not None:
                lvl_over[k] = _lvl_def(lv)
                restarts.setdefault(k, lvl_over[k][2])
        nums[nid] = (_val(n, W + "abstractNumId"), restarts, lvl_over)

    def style_num_id(style_id):
        st = styles.get(style_id)
        return _val(st, W + "pPr/" + W + "numPr/" + W + "numId") \
            if st is not None else None

    def resolve_levels(abs_id, hops=0):
        # A numStyleLink abstract borrows its levels from the list style.
        if abs_id in links and hops < 4:
            target = defines.get(links[abs_id])
            if target is None:
                nid = style_num_id(links[abs_id])
                target = nums[nid][0] if nid in nums else None
            if target is not None and target != abs_id:
                return resolve_levels(target, hops + 1)
        return abs_id, abstract.get(abs_id, {})

    def para_num(p):
        # numId and ilvl resolve independently through the style chain.
        ppr = p.find(W + "pPr")
        nid = _val(ppr, W + "numPr/" + W + "numId")
        ilvl = _val(ppr, W + "numPr/" + W + "ilvl")
        style_id, hops = _val(ppr, W + "pStyle"), 0
        while (nid is None or ilvl is None) and style_id and hops < 8:
            st = styles.get(style_id)
            if st is None:
                break
            if nid is None:
                nid = _val(st, W + "pPr/" + W + "numPr/" + W + "numId")
            if ilvl is None:
                ilvl = _val(st, W + "pPr/" + W + "numPr/" + W + "ilvl")
            style_id = _val(st, W + "basedOn")
            hops += 1
        return nid, int(ilvl or "0")

    # Word keeps one counter per abstract list; a num with a startOverride
    # restarts that list the first time the num is used.
    counters, restarted, out = {}, set(), []
    for p in paras:
        nid, ilvl = para_num(p)
        if not nid or nid == "0" or nid not in nums:
            out.append(("", None))
            continue
        abs_id, restarts, lvl_over = nums[nid]
        abs_id, levels = resolve_levels(abs_id)
        levels = {**levels, **lvl_over}
        cnt = counters.setdefault(abs_id, [None] * 9)
        ilvl = min(max(ilvl, 0), 8)
        start = levels.get(ilvl, ("decimal", "", 1, False))[2]
        if ilvl in restarts and (nid, ilvl) not in restarted:
            restarted.add((nid, ilvl))
            cnt[ilvl] = restarts[ilvl]
        else:
            cnt[ilvl] = start if cnt[ilvl] is None else cnt[ilvl] + 1
        for j in range(ilvl):
            if cnt[j] is None:
                cnt[j] = levels.get(j, ("decimal", "", 1, False))[2]
        for j in range(ilvl + 1, 9):
            cnt[j] = None
        fmt, text, _, legal = levels.get(ilvl, ("decimal", "%1.", 1, False))
        if fmt == "bullet":
            out.append(("\u2022", ilvl))
            continue

        def sub(m):
            k = int(m.group(1)) - 1
            lf = levels.get(k, ("decimal", "", 1, False))[0]
            return _fmt_num(cnt[k] if cnt[k] is not None else 1,
                            "decimal" if legal else lf)
        out.append((re.sub(r"%(\d)", sub, text).strip(), ilvl))
    return out


def numbering_labels(doc, paras):
    return [label for label, _ in numbering_info(doc, paras)]


# ── Baseline view for the reviewer ──────────────────────────────────

def _curly_share(texts):
    joined = "".join(texts)
    curly = sum(joined.count(c) for c in "“”‘’")
    straight = joined.count('"') + joined.count("'")
    return curly >= straight and curly > 0


def signature_start(texts):
    for i, t in enumerate(texts):
        if SIGNATURE_CUE.match(t or ""):
            return i
    return len(texts)


_CLOSING_CUE = re.compile(
    r"^\s*(please\b|kindly\b|if (this|the foregoing)\b|very truly|sincerely|"
    r"in witness|(schedule|exhibit|annex|appendix|attachment)\b)", re.I)


def default_insert_anchor(texts, labels, section_ends=None):
    """Where a new clause goes when no better spot is given: at the end of
    the last operative clause (its heading and any body paragraphs), before
    closing lines, schedules, a section break and the signature block."""
    sig = signature_start(texts)
    ends = section_ends or [False] * len(texts)
    numbered = [i for i in range(sig) if labels[i] and labels[i] != "\u2022"]
    anchor = numbered[-1] if numbered else None
    start = (anchor + 1) if anchor is not None else 0
    for k in range(start, sig):
        if anchor is not None and ends[anchor]:
            break
        t = (texts[k] or "").strip()
        if not t:
            if ends[k]:
                break
            continue
        if _CLOSING_CUE.match(t) or DELIVERY_CUE.match(t):
            break
        anchor = k
    if anchor is None:
        body = [i for i in range(sig) if (texts[i] or "").strip()]
        anchor = body[-1] if body else max(0, len(texts) - 1)
    return anchor


def baseline_view(source_bytes):
    """Everything the reviewer and the planner need about the baseline."""
    base = accept_all(source_bytes)
    doc = _load(base)
    paras = list_paragraphs(doc)
    texts = [para_text(p) for p in paras]
    info = numbering_info(doc, paras)
    labels = [lab for lab, _ in info]
    section_ends = [p.find(W + "pPr/" + W + "sectPr") is not None
                    for p in paras]
    return {
        "baseline": base,
        "texts": texts,
        "labels": labels,
        "levels": [lvl for _, lvl in info],
        "signature_start": signature_start(texts),
        "default_anchor": default_insert_anchor(texts, labels, section_ends),
        "curly_quotes": _curly_share(texts),
    }


def format_for_model(view):
    """The document as the reviewer sees it: one line per non-empty
    paragraph, tagged with its paragraph number and any auto-number."""
    lines = []
    sig = view["signature_start"]
    for i, (t, lab) in enumerate(zip(view["texts"], view["labels"])):
        if not (t or "").strip():
            continue
        tag = f"[¶{i}"
        if lab:
            tag += f" | auto-numbered {lab}"
        if i == sig:
            tag += " | signature block starts"
        lines.append(f"{tag}] {t.replace(chr(10), LINE_BREAK_MARK)}")
    return "\n".join(lines)


# ── Planning ─────────────────────────────────────────────────────────

def smarten_quotes(text):
    out = []
    for i, ch in enumerate(text):
        prev = text[i - 1] if i else " "
        opening = prev.isspace() or prev in "([{—–-/"
        if ch == '"':
            out.append("“" if opening else "”")
        elif ch == "'":
            out.append("‘" if opening and (i + 1 < len(text)
                                                 and text[i + 1].isalpha())
                       and not (i + 1 < len(text) and text[i + 1].isdigit())
                       else "’")
        else:
            out.append(ch)
    return "".join(out)


def _loose_pattern(s):
    out = []
    for ch in s:
        if ch.isspace():
            if not out or out[-1] != r"\s+":
                out.append(r"\s+")
        elif ch in "\"“”„":
            out.append("[\"“”„]")
        elif ch in "'‘’":
            out.append("['‘’]")
        else:
            out.append(re.escape(ch))
    return "".join(out)


def _strip_label(text, label):
    t = (text or "").strip()
    if label:
        for form in (f"({label})", label):
            if t.startswith(form + " "):
                return t[len(form):].lstrip()
    return t


DELIVERY_CUE = re.compile(
    r"^\s*(please\s+(complete|sign|execute|return|e-?mail|send|countersign)|"
    r"kindly\s+(sign|return)|return\s+(a|one)\s+(signed|executed))", re.I)
# A "new paragraph" that opens like this is really a sentence belonging to
# the clause before it; standing alone, its "this paragraph" would point
# at itself.
_CONTINUATION = re.compile(
    r"^\s*(this\s+paragraph|this\s+section|such\b|the\s+foregoing|"
    r"notwithstanding\s+the\s+foregoing|in\s+addition|for\s+the\s+avoidance"
    r"\s+of\s+doubt|provided(,)?\s+(however|that)|the\s+recipient['’]s"
    r"\s+counsel)", re.I)


def new_plan():
    return {"modify": {}, "delete": [], "insert": []}


def final_view(view, plan):
    """NCP's position as it will read once every change is accepted,
    tagged with the baseline paragraph numbers it came from."""
    texts = view["texts"]
    inserted = {}
    for ins in plan.get("insert", []):
        inserted.setdefault(ins["after"], []).append(ins["text"])
    deleted = set(plan.get("delete", []))
    lines = [f"[new clause] {t}" for t in inserted.get(-1, [])]
    for i, t in enumerate(texts):
        if i in deleted:
            lines.append(f"[¶{i} deleted]")
        else:
            body = plan.get("modify", {}).get(i, t)
            if body.strip():
                tag = f"¶{i} revised" if i in plan.get("modify", {}) else f"¶{i}"
                lines.append(f"[{tag}] {body.replace(chr(10), LINE_BREAK_MARK)}")
        lines.extend(f"[new after ¶{i}] {x.replace(chr(10), LINE_BREAK_MARK)}"
                     for x in inserted.get(i, []))
    return "\n".join(lines)


def plan_from_edits(edits, view):
    """Paragraph-keyed reviewer edits -> plan. Returns (plan, notes)."""
    texts, labels = view["texts"], view["labels"]
    n = len(texts)
    sig = view["signature_start"]
    plan, notes, pending_append = new_plan(), [], []
    for e in edits or []:
        action = (e.get("action") or "").strip().lower()
        try:
            i = int(e.get("para"))
        except (TypeError, ValueError):
            notes.append(f"Edit without a valid paragraph number skipped: "
                         f"{(e.get('text') or '')[:80]}")
            continue
        text = re.sub(r" ?⏎ ?", "\n", e.get("text") or "")
        if view.get("curly_quotes"):
            text = smarten_quotes(text)
        if action == "modify":
            if not 0 <= i < n:
                notes.append(f"¶{i} does not exist; edit skipped.")
                continue
            if i in plan["modify"]:
                notes.append(f"Two edits targeted ¶{i}; the later one was used.")
            plan["modify"][i] = _strip_label(text, labels[i])
        elif action == "delete":
            if 0 <= i < n:
                plan["delete"].append(i)
        elif action in ("insert_after", "insert"):
            body = text.strip()
            if not body:
                continue
            numbered = e.get("numbered")
            if isinstance(numbered, str) and numbered != "same_level":
                numbered = numbered.strip().lower() in ("true", "yes", "1")
            if not -1 <= i < n:
                i, numbered = view["default_anchor"], True
            if SIGNATURE_CUE.match(body):
                delivery = next((k for k in range(max(i, 0), n)
                                 if DELIVERY_CUE.match(texts[k] or "")), None)
                if delivery is not None and i >= delivery:
                    before = max([k for k in range(delivery)
                                  if (texts[k] or "").strip()] or [i])
                    notes.append(f"The acceptance block was moved above the "
                                 f"delivery instructions at ¶{delivery}.")
                    i = before
            elif i >= sig:
                notes.append(f"A new clause anchored inside the signature "
                             f"block (after ¶{i}) was moved to follow "
                             f"¶{view['default_anchor']}, the last operative "
                             f"clause.")
                i, numbered = view["default_anchor"], True
            elif i >= 0 and _CONTINUATION.match(body) and \
                    (texts[i] or "").strip():
                pending_append.append((i, body))
                notes.append(f"A sentence emitted as its own paragraph after "
                             f"¶{i} reads as part of that clause, so it was "
                             f"added to the end of ¶{i} instead.")
                continue
            plan["insert"].append({"after": i, "text": body,
                                   "numbered": numbered})
        else:
            notes.append(f"Unknown edit action '{action}' skipped.")
    for i, body in pending_append:
        base = plan["modify"].get(i, texts[i]).rstrip()
        plan["modify"][i] = f"{base} {body}"
    plan["delete"] = [i for i in sorted(set(plan["delete"]))
                      if i not in plan["modify"]]
    return plan, notes


def plan_from_legacy(edits, view):
    """Substring edits from rounds reviewed before the paragraph-keyed
    format -> plan. Returns (plan, skipped)."""
    texts = view["texts"]
    work = {}
    plan, skipped = new_plan(), []

    def cur(i):
        return work.get(i, texts[i])

    for e in edits or []:
        orig = (e.get("original") or "").strip()
        revised = e.get("revised") or ""
        if view.get("curly_quotes"):
            revised = smarten_quotes(revised)
        if not orig:
            continue
        lines = [l.strip() for l in orig.split("\n") if l.strip()]
        if len(lines) > 1:
            norm = [re.sub(r"\s+", " ", l) for l in lines]
            hit = None
            for i in range(len(texts) - len(lines) + 1):
                if [re.sub(r"\s+", " ", cur(i + k).strip())
                        for k in range(len(lines))] == norm:
                    hit = i
                    break
            if hit is None:
                skipped.append({**e, "_skip_reason":
                                "multi-paragraph text not found"})
                continue
            work[hit] = revised
            for k in range(1, len(lines)):
                plan["delete"].append(hit + k)
            continue
        pat = re.compile(_loose_pattern(orig))
        hits = [(i, m) for i in range(len(texts))
                for m in [pat.search(cur(i))] if m]
        total = sum(len(pat.findall(cur(i))) for i, _ in hits)
        if total != 1:
            skipped.append({**e, "_skip_reason":
                            f"original text found {total} times "
                            f"(needs exactly 1)"})
            continue
        i, m = hits[0]
        work[i] = cur(i)[:m.start()] + revised + cur(i)[m.end():]
    for i, t in work.items():
        plan["modify"][i] = t
    for e in edits or []:
        if (e.get("original") or "").strip():
            continue
        body = (e.get("revised") or "").strip()
        if not body:
            continue
        if view.get("curly_quotes"):
            body = smarten_quotes(body)
        anchor = (e.get("insert_after") or "").strip()
        i, numbered = view["default_anchor"], True
        if anchor:
            pat = re.compile(_loose_pattern(anchor))
            found = [k for k in range(len(texts)) if pat.search(cur(k))]
            if found and found[0] < view["signature_start"]:
                i = found[0]
        plan["insert"].append({"after": i, "text": body,
                               "numbered": numbered})
    plan["delete"] = [i for i in sorted(set(plan["delete"]))
                      if i not in plan["modify"]]
    return plan, skipped


# ── Word-level comparison ────────────────────────────────────────────

_TOKEN = re.compile(r"\s+|\w+|[^\w\s]", re.UNICODE)
_QMAP = str.maketrans({"“": '"', "”": '"', "„": '"',
                       "‘": "'", "’": "'", "‑": "-"})


def _norm_tok(t):
    return " " if t.isspace() else t.translate(_QMAP)


def norm_text(s):
    return "".join(_norm_tok(t) for t in _TOKEN.findall(s or "")).strip()


def _is_glue(tokens):
    words = [t for t in tokens if not t.isspace() and re.match(r"\w", t)]
    return len(words) <= 2 and sum(len(t) for t in tokens) <= 14


def diff_ops(old, new):
    """Word-level edit script in character coordinates.

    Returns (ops, merged) where ops are (tag, a0, a1, b0, b1) with tag in
    equal/delete/insert/replace, and merged is the paragraph text after
    accepting (equal spans keep the counterparty's exact characters)."""
    ot, nt = _TOKEN.findall(old), _TOKEN.findall(new)
    on, nn = [_norm_tok(t) for t in ot], [_norm_tok(t) for t in nt]
    sm = difflib.SequenceMatcher(None, on, nn, autojunk=False)
    ops = sm.get_opcodes()
    # Absorb short unchanged glue ("the", "of the") sitting between two
    # changes so the banker reads whole phrases, not confetti.
    changed = True
    while changed:
        changed, out, k = False, [], 0
        while k < len(ops):
            if (out and k + 1 < len(ops) and ops[k][0] == "equal"
                    and out[-1][0] != "equal" and ops[k + 1][0] != "equal"
                    and _is_glue(ot[ops[k][1]:ops[k][2]])):
                prev, nxt = out.pop(), ops[k + 1]
                out.append(("replace", prev[1], nxt[2], prev[3], nxt[4]))
                k += 2
                changed = True
                continue
            out.append(ops[k])
            k += 1
        ops = out
    # Heavily rewritten paragraph: one deletion and one insertion for the
    # changed region (the rulebook's fix for fragmented compare output).
    non_equal = [o for o in ops if o[0] != "equal"]
    if len(non_equal) > 1 and sm.ratio() < 0.45:
        head = ops[0] if ops[0][0] == "equal" else None
        tail = ops[-1] if ops[-1][0] == "equal" and len(ops) > 1 else None
        a0 = head[2] if head else 0
        b0 = head[4] if head else 0
        a1 = tail[1] if tail else len(ot)
        b1 = tail[3] if tail else len(nt)
        ops = ([head] if head else []) + [("replace", a0, a1, b0, b1)] + \
              ([tail] if tail else [])

    def offsets(toks):
        acc, pos = [], 0
        for t in toks:
            acc.append(pos)
            pos += len(t)
        acc.append(pos)
        return acc
    oo, no = offsets(ot), offsets(nt)
    char_ops, merged = [], []
    for tag, i1, i2, j1, j2 in ops:
        a0, a1, b0, b1 = oo[i1], oo[i2], no[j1], no[j2]
        if tag == "replace" and a0 == a1:
            tag = "insert"
        elif tag == "replace" and b0 == b1:
            tag = "delete"
        char_ops.append((tag, a0, a1, b0, b1))
        merged.append(old[a0:a1] if tag == "equal" else new[b0:b1])
    return char_ops, "".join(merged)


# ── Writing tracked changes into the counterparty's runs ────────────

class _Ctx:
    def __init__(self, author, date):
        self.author, self.date, self.next_id = author, date, ID_START

    def rev(self, kind):
        el = _el(kind)
        el.set(W + "id", str(self.next_id))
        el.set(W + "author", self.author)
        el.set(W + "date", self.date)
        self.next_id += 1
        return el


def _collect_runs(p):
    """[(run, start, end, top)] in reading order; top is the paragraph's
    direct child that holds the run (the run itself, or e.g. a hyperlink)."""
    out, pos = [], [0]

    def walk(el, top):
        for ch in el:
            if not _is_tag(ch):
                continue
            if ch.tag == W + "r":
                n = len(_run_text(ch))
                out.append((ch, pos[0], pos[0] + n,
                            ch if top is None else top))
                pos[0] += n
            elif ch.tag in _CONTAINERS or ch.tag in _INS:
                walk(ch, ch if top is None else top)
    walk(p, None)
    return out


_VOLATILE_FIELDS = {"DATE", "TIME", "PRINTDATE", "SAVEDATE", "CREATEDATE",
                    "EDITTIME", "REVNUM", "SEQ", "AUTONUM", "AUTONUMLGL",
                    "AUTONUMOUT", "LISTNUM"}
_OBJECT_TAGS = {W + "drawing", W + "pict", W + "object", MC + "AlternateContent",
                W + "txbxContent"}
# Zero-width run content that still prints (and so must go when its
# clause is deleted).
_PRINTING_ATOMS = {W + "footnoteReference", W + "endnoteReference", W + "sym",
                   W + "drawing", W + "pict", W + "object",
                   MC + "AlternateContent", W + "commentReference"}


def _outside_objects(el, p):
    a = el.getparent()
    while a is not None and a is not p:
        if a.tag in _OBJECT_TAGS:
            return False
        a = a.getparent()
    return True


def _collapse_fields(p):
    """Turn volatile fields (DATE, SEQ and the like) that sit wholly inside
    this paragraph's own runs into their displayed text (rulebook: they
    otherwise surface as spurious changes). Other fields, and anything
    inside text boxes or drawings, keep working."""
    for fs in [x for x in p.iter(W + "fldSimple") if _outside_objects(x, p)]:
        code = (fs.get(W + "instr") or "").split()
        if code and code[0].upper() in _VOLATILE_FIELDS:
            _unwrap(fs)
    stack, complete = [], []
    for r, _, _, _ in _collect_runs(p):
        for ch in r:
            if not _is_tag(ch):
                continue
            if ch.tag == W + "fldChar":
                kind = ch.get(W + "fldCharType")
                if kind == "begin":
                    stack.append({"runs": [r], "code": ""})
                elif kind == "separate" and stack:
                    stack[-1]["runs"].append(r)
                elif kind == "end" and stack:
                    f = stack.pop()
                    f["runs"].append(r)
                    if not stack:
                        complete.append(f)
            elif ch.tag == W + "instrText" and stack:
                stack[-1]["code"] += ch.text or ""
                stack[-1]["runs"].append(r)
    for f in complete:
        code = f["code"].split()
        if not code or code[0].upper() not in _VOLATILE_FIELDS:
            continue
        for r in f["runs"]:
            for ch in list(r):
                if _is_tag(ch) and ch.tag in (W + "fldChar", W + "instrText"):
                    r.remove(ch)
            if r.getparent() is not None and \
                    not [c for c in r if _is_tag(c) and c.tag != W + "rPr"]:
                r.getparent().remove(r)


def _clean_rpr(rpr):
    if rpr is None:
        return None
    rpr = copy.deepcopy(rpr)
    for ch in list(rpr):
        if not _is_tag(ch):
            continue
        if ch.tag in (W + "ins", W + "del", W + "moveFrom", W + "moveTo",
                      W + "rPrChange"):
            rpr.remove(ch)
        elif ch.tag == W + "rStyle" and (ch.get(W + "val") or "").lower() in (
                "hyperlink", "followedhyperlink"):
            rpr.remove(ch)
    return rpr


def _text_run(text, rpr):
    r = _el("r")
    if rpr is not None:
        r.append(copy.deepcopy(rpr))
    for piece in re.split(r"(\t|\n)", text):
        if not piece:
            continue
        if piece == "\t":
            r.append(_el("tab"))
        elif piece == "\n":
            r.append(_el("br"))
        else:
            t = _el("t")
            t.text = piece
            t.set(XML_SPACE, "preserve")
            r.append(t)
    return r


def _slice_run(r, rs, slices):
    rpr = r.find(W + "rPr")
    outs = []
    for _ in slices:
        nr = _el("r")
        for k, v in r.attrib.items():
            nr.set(k, v)
        if rpr is not None:
            nr.append(copy.deepcopy(rpr))
        outs.append(nr)
    pos = rs
    for child, txt in _run_atoms(r):
        n = len(txt)
        if n == 0:
            k = next((j for j, (s, e) in enumerate(slices) if s <= pos < e),
                     len(slices) - 1)
            outs[k].append(copy.deepcopy(child))
            continue
        for k, (s, e) in enumerate(slices):
            lo, hi = max(pos, s), min(pos + n, e)
            if lo >= hi:
                continue
            if child.tag in (W + "t", W + "delText"):
                t = _el("t")
                t.text = txt[lo - pos:hi - pos]
                t.set(XML_SPACE, "preserve")
                outs[k].append(t)
            else:
                outs[k].append(copy.deepcopy(child))
        pos += n
    return outs


def _as_deleted(r):
    for ch in list(r):
        if not _is_tag(ch):
            continue
        if ch.tag == W + "t":
            d = _el("delText")
            d.text = ch.text
            d.set(XML_SPACE, "preserve")
            r.replace(ch, d)
        elif ch.tag == W + "instrText":
            d = _el("delInstrText")
            d.text = ch.text
            r.replace(ch, d)
    return r


def _rpr_key(r):
    rpr = r.find(W + "rPr")
    if rpr is None:
        return b""
    from lxml import etree
    return etree.tostring(_clean_rpr(rpr))


def _apply_ops(p, old, ops, ctx, delete_all=False):
    """Rewrite paragraph p in place so it carries the edit script. With
    delete_all, every run goes, including zero-width ones (footnote marks,
    symbols, images)."""
    runs = _collect_runs(p)
    cuts = set()
    deleted, inserts = [], []
    for tag, a0, a1, b0, b1 in ops:
        cuts.update((a0, a1))
        if tag in ("delete", "replace") and a1 > a0:
            deleted.append((a0, a1))
        if tag in ("insert", "replace") and b1 > b0:
            inserts.append((a1 if tag == "replace" else a0, (b0, b1)))

    def is_del(s, e):
        return any(a0 <= s and e <= a1 for a0, a1 in deleted)

    slices_of = {}
    for r, rs, re_, top in runs:
        if re_ == rs:
            printing = any(_is_tag(c) and c.tag in _PRINTING_ATOMS for c in r)
            if delete_all or (printing and any(a0 < rs < a1
                                               for a0, a1 in deleted)):
                d = ctx.rev("del")
                r.getparent().replace(r, d)
                d.append(_as_deleted(r))
            continue
        bounds = sorted({rs, re_} | {c for c in cuts if rs < c < re_})
        spans = list(zip(bounds, bounds[1:]))
        if len(spans) == 1 and not is_del(*spans[0]):
            slices_of[id(r)] = [(rs, re_, r)]
            continue
        built = _slice_run(r, rs, spans)
        elems = []
        for (s, e), nr in zip(spans, built):
            if is_del(s, e):
                d = ctx.rev("del")
                d.append(_as_deleted(nr))
                elems.append((s, e, d))
            else:
                elems.append((s, e, nr))
        parent = r.getparent()
        at = parent.index(r)
        parent.remove(r)
        for k, (_, _, el) in enumerate(elems):
            parent.insert(at + k, el)
        slices_of[id(r)] = elems
    return runs, slices_of, inserts


def _place_insertions(p, old, new, runs, slices_of, inserts, ctx):
    text_runs = [x for x in runs if x[2] > x[1]]
    weights = {}
    for r, rs, re_, _ in text_runs:
        k = _rpr_key(r)
        weights[k] = weights.get(k, 0) + (re_ - rs)
    dominant = max(weights, key=weights.get) if weights else b""
    span_of_top = {}
    for r, rs, re_, top in text_runs:
        lo, hi = span_of_top.get(id(top), (rs, re_))
        span_of_top[id(top)] = (min(lo, rs), max(hi, re_))
    last_after = {}

    def in_link(top, r):
        return top is not r and top.tag == W + "hyperlink"

    for x, (b0, b1) in inserts:
        text = new[b0:b1]
        prev = next(((r, rs, re_, t) for r, rs, re_, t in text_runs
                     if rs < x <= re_), None)
        nxt = next(((r, rs, re_, t) for r, rs, re_, t in text_runs
                    if rs <= x < re_), None)
        if prev and nxt:
            if in_link(prev[3], prev[0]) and not in_link(nxt[3], nxt[0]):
                choice = (nxt, "before")
            elif in_link(nxt[3], nxt[0]) and not in_link(prev[3], prev[0]):
                choice = (prev, "after")
            elif _rpr_key(prev[0]) == dominant or _rpr_key(nxt[0]) != dominant:
                choice = (prev, "after")
            else:
                choice = (nxt, "before")
        elif prev:
            choice = (prev, "after")
        elif nxt:
            choice = (nxt, "before")
        else:
            choice = (None, "append")
        (anchor, side) = choice
        if anchor is None:
            rpr = _clean_rpr(p.find(W + "pPr/" + W + "rPr"))
            ins = ctx.rev("ins")
            ins.append(_text_run(text, rpr))
            key = id(p)
            if key in last_after:
                last_after[key].addnext(ins)
            else:
                p.append(ins)
            last_after[key] = ins
            continue
        r, rs, re_, top = anchor
        ins = ctx.rev("ins")
        ins.append(_text_run(text, _clean_rpr(r.find(W + "rPr"))))
        elems = slices_of[id(r)]
        if side == "after":
            target = next(el for s, e, el in elems if e == x)
            if top is not r and top.tag == W + "hyperlink" and \
                    span_of_top[id(top)][1] == x:
                target = top
            if id(target) in last_after:
                last_after[id(target)].addnext(ins)
            else:
                target.addnext(ins)
            last_after[id(target)] = ins
        else:
            target = next(el for s, e, el in elems if s == x)
            if top is not r and top.tag == W + "hyperlink" and \
                    span_of_top[id(top)][0] == x:
                target = top
            target.addprevious(ins)


def _mark_paragraph(p, ctx, kind):
    ppr = p.find(W + "pPr")
    if ppr is None:
        ppr = _el("pPr")
        p.insert(0, ppr)
    rpr = ppr.find(W + "rPr")
    if rpr is None:
        rpr = _el("rPr")
        before = None
        for tag in ("sectPr", "pPrChange"):
            before = ppr.find(W + tag)
            if before is not None:
                break
        if before is not None:
            before.addprevious(rpr)
        else:
            ppr.append(rpr)
    rpr.insert(0, ctx.rev(kind))


def _caption_rpr(p):
    """(caption_rPr, body_rPr) when the clause opens with a bold/underlined
    caption like 'Confidentiality.' and continues in body text."""
    runs = [x for x in _collect_runs(p) if x[2] > x[1]]
    if len(runs) < 2:
        return None
    first = runs[0][0]
    rpr = first.find(W + "rPr")
    if rpr is None:
        return None
    b = rpr.find(W + "b")
    u = rpr.find(W + "u")
    emphasized = (b is not None and b.get(W + "val") not in ("0", "false")) or \
                 (u is not None and u.get(W + "val") not in (None, "none"))
    text = _run_text(first).strip()
    if not emphasized or not text or len(text) > 80 or text[-1] not in ".:":
        return None
    body = next((x[0] for x in runs[1:] if _rpr_key(x[0]) != _rpr_key(first)),
                None)
    if body is None:
        return None
    return _clean_rpr(rpr), _clean_rpr(body.find(W + "rPr"))


def _new_paragraph(template, text, ctx):
    p = _el("p")
    tppr = template.find(W + "pPr") if template is not None else None
    if tppr is not None:
        ppr = copy.deepcopy(tppr)
        # Never inherit a neighbor's section break, fill-in underline
        # border, forced page break or text frame.
        for tag in ("sectPr", "pPrChange", "pBdr", "pageBreakBefore",
                    "framePr"):
            for el in ppr.findall(W + tag):
                ppr.remove(el)
        prpr = ppr.find(W + "rPr")
        if prpr is not None:
            for ch in list(prpr):
                if _is_tag(ch) and ch.tag in (W + "ins", W + "del",
                                              W + "moveFrom", W + "moveTo",
                                              W + "rPrChange"):
                    prpr.remove(ch)
        p.append(ppr)
    _mark_paragraph(p, ctx, "ins")
    runs = [x for x in _collect_runs(template) if x[2] > x[1]] \
        if template is not None else []
    weights = {}
    for r, rs, re_, _ in runs:
        k = _rpr_key(r)
        weights.setdefault(k, [0, r])[0] += re_ - rs
    body_rpr = None
    if weights:
        body_rpr = _clean_rpr(max(weights.values(), key=lambda v: v[0])[1]
                              .find(W + "rPr"))
    elif tppr is not None:
        body_rpr = _clean_rpr(tppr.find(W + "rPr"))
    ins = ctx.rev("ins")
    cap = _caption_rpr(template) if template is not None else None
    m = re.match(r"^([A-Z][^.:\n]{1,78}[.:])(\s+)(.+)$", text, re.S)
    if cap and m:
        ins.append(_text_run(m.group(1), cap[0]))
        ins.append(_text_run(m.group(2) + m.group(3), cap[1]))
    else:
        ins.append(_text_run(text, body_rpr))
    p.append(ins)
    return p


def _has_num(labels, i):
    return bool(labels[i]) and labels[i] != "\u2022"


def _template_for(paras, labels, levels, after, numbered):
    """The paragraph whose properties a new paragraph copies.

    numbered True: a new top-level clause, copying the last top-level
    numbered clause (never a sub-item or a heading's body). "same_level":
    another item at the anchor's own numbering level, e.g. (c) after (b).
    None: follow the anchor. False: unnumbered, laid out like the anchor."""
    if after < 0:
        return paras[0] if paras else None
    anchor = paras[after]
    if numbered in (False, None):
        return anchor
    numbered_before = [k for k in range(after, -1, -1) if _has_num(labels, k)]
    if numbered == "same_level":
        return paras[numbered_before[0]] if numbered_before else anchor
    top = [k for k in numbered_before if levels[k] == 0]
    if top:
        return paras[top[0]]
    if numbered_before:
        return paras[numbered_before[0]]
    return anchor


def _strip_numbering(p):
    ppr = p.find(W + "pPr")
    if ppr is not None:
        for el in ppr.findall(W + "numPr"):
            ppr.remove(el)


# ── Build + audit ────────────────────────────────────────────────────

def _schema_problems(body, author):
    problems = []
    ids = []
    for el in body.iter(W + "ins", W + "del", W + "moveFrom", W + "moveTo"):
        if el.get(W + "author") != author:
            problems.append(f"revision authored by {el.get(W + 'author')!r}")
        try:
            ids.append(int(el.get(W + "id")))
        except (TypeError, ValueError):
            problems.append("revision without a numeric id")
        holder = el.getparent()
        if holder is not None and holder.tag == W + "rPr" and \
                [c for c in holder if _is_tag(c)][0] is not el:
            problems.append("revision mark is not the first child of rPr")
    if len(ids) != len(set(ids)):
        problems.append("duplicate revision ids")
    if ids and min(ids) < ID_START:
        problems.append(f"revision ids start below {ID_START}")
    for d in body.iter(W + "del"):
        if any(True for _ in d.iter(W + "t")):
            problems.append("w:t inside a deletion (must be w:delText)")
            break
    for t in body.iter(W + "delText"):
        a = t.getparent()
        while a is not None and a.tag not in (W + "del", W + "moveFrom"):
            a = a.getparent()
        if a is None:
            problems.append("w:delText outside a deletion")
            break
    return problems


def build(source_bytes, plan, author=AUTHOR, when=None):
    """Build the redline and the clean copy from the counterparty's file.

    Returns (redline_bytes, clean_bytes, report). report["ok"] is True only
    when every audit passed; report["expected"] is NCP's position as
    paragraph texts."""
    from lxml import etree

    base_bytes = accept_all(source_bytes)
    doc = _load(base_bytes)
    paras = list_paragraphs(doc)
    info = numbering_info(doc, paras)
    labels = [lab for lab, _ in info]
    levels = [lvl for _, lvl in info]
    base_texts = [para_text(p) for p in paras]
    n = len(paras)
    date = when or datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    ctx = _Ctx(author, date)
    notes, applied = [], []
    expected = list(base_texts)
    changed = set()

    for i, new in sorted(plan.get("modify", {}).items()):
        if not 0 <= i < n:
            notes.append(f"¶{i} does not exist; edit skipped.")
            continue
        if norm_text(new) == norm_text(base_texts[i]):
            continue
        p = paras[i]
        _collapse_fields(p)
        old = para_text(p)
        ops, merged = diff_ops(old, new)
        runs, slices_of, inserts = _apply_ops(p, old, ops, ctx)
        _place_insertions(p, old, new, runs, slices_of, inserts, ctx)
        expected[i] = merged
        changed.add(i)
        applied.append(("modify", i))

    for i in plan.get("delete", []):
        if not 0 <= i < n or i in changed:
            continue
        p = paras[i]
        _collapse_fields(p)
        old = para_text(p)
        _apply_ops(p, old, [("delete", 0, len(old), 0, 0)] if old else [],
                   ctx, delete_all=True)
        nxt = p.getnext()
        while nxt is not None and not _is_tag(nxt):
            nxt = nxt.getnext()
        # A paragraph mark that carries a section break stays, or accepting
        # the deletion would take the page layout with it.
        holds_section = p.find(W + "pPr/" + W + "sectPr") is not None
        if nxt is not None and nxt.tag == W + "p" and not holds_section:
            _mark_paragraph(p, ctx, "del")
            expected[i] = None
        else:
            expected[i] = ""
        changed.add(i)
        applied.append(("delete", i))

    inserted = {}
    last_new = {}
    for ins in plan.get("insert", []):
        after = ins.get("after", -1)
        if not -1 <= after < n:
            notes.append("A new clause had no valid position; skipped.")
            continue
        text = (ins.get("text") or "").strip()
        if not text:
            continue
        template = _template_for(paras, labels, levels, after,
                                 ins.get("numbered"))
        newp = _new_paragraph(template, text, ctx)
        if ins.get("numbered") is False:
            _strip_numbering(newp)
        if after in last_new:
            last_new[after].addnext(newp)
        elif after == -1:
            paras[0].addprevious(newp)
        else:
            paras[after].addnext(newp)
        last_new[after] = newp
        inserted.setdefault(after, []).append(text)
        applied.append(("insert", after))

    # Like pressing Enter at the end of a paragraph that ends a section:
    # the section break moves to the last new paragraph, so new clauses
    # print in the anchor's section rather than the next one.
    for after, newp in last_new.items():
        if after < 0:
            continue
        sect = paras[after].find(W + "pPr/" + W + "sectPr")
        if sect is not None:
            sect.getparent().remove(sect)
            newp.find(W + "pPr").append(sect)
            changed.add(after)
            notes.append(f"¶{after} ends a document section; its section "
                         f"break now follows the new text inserted after it.")

    exp = list(inserted.get(-1, []))
    for i in range(n):
        if expected[i] is not None:
            exp.append(expected[i])
        exp.extend(inserted.get(i, []))

    redline = _save(doc)

    # ── Audits (on freshly parsed files, never on in-memory state) ──
    audits, problems = {}, []
    rdoc = _load(redline)
    rparas = list_paragraphs(rdoc)
    rej = doc_view(rparas, "reject")
    acc = doc_view(rparas, "accept")
    audits["reject"] = rej == base_texts
    if not audits["reject"]:
        problems.append("Rejecting all changes does not reproduce the "
                        "counterparty's text.")
    audits["accept"] = acc == exp
    if not audits["accept"]:
        problems.append("Accepting all changes does not equal NCP's "
                        "position.")
    clean = accept_all(redline)
    cdoc = _load(clean)
    ctexts = [para_text(p) for p in list_paragraphs(cdoc)]
    leftover = story_revisions(clean)
    audits["clean"] = ctexts == exp and leftover == 0
    if not audits["clean"]:
        problems.append("The clean copy does not match NCP's position.")
    bdoc = _load(base_bytes)
    bparas = list_paragraphs(bdoc)
    kept = [p for p in rparas if not (_mark(p) == "ins"
                                      and not para_text(p, "reject"))]
    untouched_ok = len(kept) == len(bparas)
    if untouched_ok:
        for i, (bp, rp) in enumerate(zip(bparas, kept)):
            if i in changed:
                continue
            if etree.tostring(bp) != etree.tostring(rp):
                untouched_ok = False
                problems.append(f"¶{i} changed although NCP did not edit it.")
                break
    else:
        problems.append("Paragraph structure drifted from the baseline.")
    audits["untouched"] = untouched_ok
    schema = _schema_problems(rdoc.element.body, author)
    audits["schema"] = not schema
    problems.extend(sorted(set(schema)))
    audits["count"] = len(acc) == len(exp)

    dash = [t for t in [*(plan.get("modify", {}).values()),
                        *(x.get("text", "") for x in plan.get("insert", []))]
            if "—" in t or "–" in t]
    report = {
        "ok": all(audits.values()),
        "audits": audits,
        "problems": problems,
        "notes": notes,
        "applied": applied,
        "changes": ctx.next_id - ID_START,
        "author": author,
        "first_id": ID_START,
        "expected": exp,
        "baseline": base_texts,
        "style_warnings": (["NCP text contains an em or en dash; the "
                            "rulebook says to use commas, periods or "
                            "parentheses instead."] if dash else []),
    }
    return redline, clean, report


# ── Counter-round support ────────────────────────────────────────────

def incoming_changes(upload_bytes, limit=40):
    """The counterparty's live tracked changes, per paragraph, as
    (before, after) text pairs."""
    doc = _load(upload_bytes)
    out = []
    for p in list_paragraphs(doc):
        has_rev = any(True for _ in p.iter(W + "ins", W + "del",
                                           W + "moveFrom", W + "moveTo"))
        if not has_rev:
            continue
        before, after = para_text(p, "reject"), para_text(p, "accept")
        if before != after or _mark(p):
            authors = sorted({el.get(W + "author") or "?"
                              for el in p.iter(W + "ins", W + "del")})
            out.append({"before": before, "after": after,
                        "authors": authors, "mark": _mark(p)})
        if len(out) >= limit:
            break
    return out


def untracked_changes(prev_texts, new_texts, limit=40):
    """Paragraph-level differences between NCP's last position and the
    counterparty's new file (for files sent without live markup)."""
    a = [norm_text(t) for t in prev_texts]
    b = [norm_text(t) for t in new_texts]
    sm = difflib.SequenceMatcher(None, a, b, autojunk=False)
    out = []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            continue
        out.append({"kind": tag,
                    "ncp_last": [prev_texts[k] for k in range(i1, i2)
                                 if prev_texts[k].strip()],
                    "theirs": [new_texts[k] for k in range(j1, j2)
                               if new_texts[k].strip()],
                    "paras": list(range(j1, j2))})
        if len(out) >= limit:
            break
    return [o for o in out if o["ncp_last"] or o["theirs"]]


def docx_from_text(text):
    """A plain .docx rebuilt from extracted PDF text, so a PDF upload still
    gets a genuine tracked-changes redline (layout is not preserved)."""
    import docx
    lines = (text or "").replace("\r", "").split("\n")
    paras, buf = [], []
    for line in lines:
        s = line.strip()
        if not s:
            if buf:
                paras.append(" ".join(buf))
                buf = []
            continue
        if buf and (re.match(r"^(\(?[0-9a-zA-Z]{1,3}[.)]|[•\-])\s", s)
                    or re.search(r"[.:;]$", buf[-1])):
            paras.append(" ".join(buf))
            buf = []
        buf.append(s)
    if buf:
        paras.append(" ".join(buf))
    d = docx.Document()
    for t in paras:
        d.add_paragraph(t)
    return _save(d)
